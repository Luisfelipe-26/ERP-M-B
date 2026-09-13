"""Reporte de ejecución: apropiado vs consumido, agrupaciones, drill-down y export."""
import datetime as dt
from io import BytesIO

import openpyxl
import pytest

import models
from conftest import ANIO, mover, presupuestar
from routers.contabilidad import (ejecucion_presupuestaria,
                                  ejecucion_presupuestaria_detalle,
                                  exportar_ejecucion_presupuestaria)


@pytest.fixture
def dimensiones(db):
    db.add_all([
        models.Departamento(id=7, codigo="AGR", nombre="Agrícola", activo=True),
        models.Departamento(id=9, codigo="MAQ", nombre="Maquinaria", activo=True),
    ])
    db.commit()


@pytest.fixture
def compra_completa(db, cuenta, proveedor, dimensiones):
    """Una OC facturada y pagada en el depto 7, y otra OC solo comprometida en el 9."""
    presupuestar(db, cuenta.id, 100_000, meses=2, departamento_id=7)
    presupuestar(db, cuenta.id, 60_000, meses=2, departamento_id=9)

    db.add(models.OrdenCompra(oc_id="OC-001", fecha=dt.datetime(ANIO, 1, 10),
                              proveedor=proveedor.nombre, proveedor_id=proveedor.id,
                              estado="Recibida", total_estimado=50_000))
    db.add(models.CuentaPorPagar(numero="CXP-001", proveedor_id=proveedor.id,
                                 oc_id="OC-001", fecha_factura=dt.date(ANIO, 1, 20),
                                 subtotal=50_000, itbis=9_000, total=59_000,
                                 saldo_pendiente=0, estado="pagada"))
    db.commit()

    mover(db, cuenta.id, "COMPROMISO", 50_000, departamento_id=7, origen_tipo="OC", origen_id="OC-001")
    mover(db, cuenta.id, "LIBERACION", -50_000, departamento_id=7, origen_tipo="CXP", origen_id="CXP-001")
    mover(db, cuenta.id, "DEVENGADO", 50_000, departamento_id=7, origen_tipo="CXP", origen_id="CXP-001")
    mover(db, cuenta.id, "PAGADO", 50_000, departamento_id=7, origen_tipo="PAGO", origen_id="PAG-001")
    mover(db, cuenta.id, "COMPROMISO", 20_000, departamento_id=9, origen_tipo="OC", origen_id="OC-002")


def test_agrupar_por_linea_abre_una_fila_por_dimension(db, user, cuenta, compra_completa):
    filas = ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="linea", db=db, user=user)
    por_depto = {f["departamento_nombre"]: f for f in filas}

    assert len(filas) == 2
    assert por_depto["Agrícola"]["apropiado"] == 100_000
    assert por_depto["Agrícola"]["comprometido"] == 0, "la factura liberó el compromiso"
    assert por_depto["Agrícola"]["devengado"] == 50_000
    assert por_depto["Agrícola"]["disponible"] == 50_000
    assert por_depto["Agrícola"]["pct_ejecucion"] == 50.0
    assert por_depto["Maquinaria"]["comprometido"] == 20_000
    assert por_depto["Maquinaria"]["disponible"] == 40_000


def test_agrupar_por_cuenta_consolida(db, user, cuenta, compra_completa):
    filas = ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="cuenta", db=db, user=user)

    assert len(filas) == 1
    assert filas[0]["apropiado"] == 160_000
    assert filas[0]["disponible"] == 90_000


def test_los_totales_coinciden_entre_agrupaciones(db, user, cuenta, compra_completa):
    por_linea = ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="linea", db=db, user=user)
    por_cuenta = ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="cuenta", db=db, user=user)

    for campo in ("apropiado", "comprometido", "devengado", "disponible"):
        assert sum(f[campo] for f in por_linea) == pytest.approx(
            sum(f[campo] for f in por_cuenta)), campo


def test_ytd_acumula_lo_apropiado_sin_cambiar_el_consumo(db, user, cuenta, compra_completa):
    febrero = ejecucion_presupuestaria(anio=ANIO, mes=2, agrupar="linea", db=db, user=user)
    agricola = next(f for f in febrero if f["departamento_id"] == 7)

    assert agricola["apropiado"] == 200_000
    assert agricola["devengado"] == 50_000
    assert agricola["disponible"] == 150_000


def test_marca_el_consumo_sin_presupuesto(db, user, cuenta, dimensiones):
    """Gasto sobre una combinación que nadie presupuestó: debe verse, no desaparecer."""
    mover(db, cuenta.id, "DEVENGADO", 15_000, departamento_id=9, origen_tipo="CXP", origen_id="CXP-X")

    filas = ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="linea", db=db, user=user)

    assert len(filas) == 1
    assert filas[0]["sin_presupuesto"] is True
    assert filas[0]["disponible"] == -15_000


def test_drill_down_lista_los_documentos_de_la_linea(db, user, cuenta, compra_completa):
    det = ejecucion_presupuestaria_detalle(anio=ANIO, mes=1, cuenta_id=cuenta.id,
                                           departamento_id=7, db=db, user=user)

    assert det["total"] == 4
    assert det["devengado"] == 50_000
    assert det["comprometido"] == 0

    oc = next(i for i in det["items"] if i["origen_tipo"] == "OC")
    assert oc["proveedor"] == "Agroquímica SA"
    assert oc["estado_documento"] == "Recibida"

    factura = next(i for i in det["items"] if i["tipo"] == "DEVENGADO")
    assert factura["oc_vinculada"] == "OC-001"


def test_drill_down_sin_movimientos(db, user, cuenta):
    det = ejecucion_presupuestaria_detalle(anio=ANIO, mes=1, cuenta_id=cuenta.id,
                                           db=db, user=user)
    assert det == {"total": 0, "items": []}


def test_export_produce_un_xlsx_con_totales(db, user, cuenta, compra_completa):
    resp = exportar_ejecucion_presupuestaria(anio=ANIO, mes=1, agrupar="linea",
                                             db=db, user=user)

    assert resp.body[:2] == b"PK", "un .xlsx es un zip"
    assert "ejecucion_presupuestaria_2026_m1.xlsx" in resp.headers["content-disposition"]

    ws = openpyxl.load_workbook(BytesIO(resp.body)).active
    assert ws["A1"].value.startswith("Ejecución presupuestaria 2026 — acumulado a enero")
    assert ws["A3"].value == "Cuenta"
    assert ws["I3"].value == "Disponible"
    assert ws.max_row == 6, "cabecera + 2 líneas + total"
    assert str(ws["A6"].value) == "TOTAL"
    assert ws.freeze_panes == "A4"
